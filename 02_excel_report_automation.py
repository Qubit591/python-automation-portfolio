"""
Excel Report Automation
Reads raw sales data CSV, generates a formatted Excel report with charts.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from datetime import datetime
import sys


def load_data(filepath):
    df = pd.read_csv(filepath, parse_dates=["date"])
    df["month"] = df["date"].dt.to_period("M").astype(str)
    return df


def build_summary(df):
    return df.groupby("month").agg(
        total_sales=("amount", "sum"),
        nb_orders=("order_id", "count"),
        avg_order=("amount", "mean")
    ).reset_index()


def style_header(ws, row, cols):
    fill = PatternFill("solid", fgColor="5B21B6")
    font = Font(bold=True, color="FFFFFF", size=11)
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def add_borders(ws, min_row, max_row, min_col, max_col):
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border


def generate_report(input_csv, output_path=None):
    df = load_data(input_csv)
    summary = build_summary(df)

    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Summary"

    # Title
    ws["A1"] = f"Sales Report — Generated {datetime.now().strftime('%d/%m/%Y')}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 30

    # Headers
    headers = ["Month", "Total Sales (€)", "Orders", "Avg Order (€)"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=3, column=col, value=h)
    style_header(ws, 3, len(headers))

    # Data
    for i, row in summary.iterrows():
        r = i + 4
        ws.cell(r, 1, row["month"])
        ws.cell(r, 2, round(row["total_sales"], 2))
        ws.cell(r, 3, row["nb_orders"])
        ws.cell(r, 4, round(row["avg_order"], 2))
        if i % 2 == 0:
            for col in range(1, 5):
                ws.cell(r, col).fill = PatternFill("solid", fgColor="F3F0FF")

    last_row = len(summary) + 3
    add_borders(ws, 3, last_row, 1, 4)

    # Column widths
    for col in range(1, 5):
        ws.column_dimensions[get_column_letter(col)].width = 20

    # Chart
    chart = BarChart()
    chart.title = "Monthly Sales"
    chart.y_axis.title = "€"
    chart.x_axis.title = "Month"
    chart.style = 10
    chart.width = 20
    chart.height = 12

    data_ref = Reference(ws, min_col=2, min_row=3, max_row=last_row)
    cats_ref = Reference(ws, min_col=1, min_row=4, max_row=last_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws.add_chart(chart, f"F3")

    output = output_path or f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(output)
    print(f"✅ Report saved: {output}")
    return output


if __name__ == "__main__":
    csv_file = sys.argv[1] if len(sys.argv) > 1 else "sales_data.csv"
    generate_report(csv_file)
